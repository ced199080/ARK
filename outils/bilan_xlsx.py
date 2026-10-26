#!/usr/bin/env python3
"""Bilan matériaux en Excel (donnees/07_bilan_materiaux.xlsx), depuis donnees/synthese.json produit par forteresse.py.

Usage : python3 outils/bilan_xlsx.py   (après python3 outils/forteresse.py)
Onglets : Synthèse (structures et matériaux par poste, ressources brutes), Détail (une ligne par pièce : nombre et coût
unitaire modifiables, tout se recalcule), Par plan (pièces de chaque plan, options comprises).
"""
import json
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import forteresse as fz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

S = json.load(open(os.path.join(fz.RACINE, 'donnees', 'synthese.json'), encoding='utf-8'))
RES = [('MI', 'Lingots de métal'), ('PC', 'Pâte cimentée'), ('EL', 'Électroniques'), ('PO', 'Polymère'),
       ('CR', 'Cristal'), ('EM', 'Élément'), ('PN', 'Perles noires'), ('HU', 'Huile'), ('PE', 'Poudre étincelante'),
       ('CU', 'Cuir'), ('FI', 'Fibre'), ('CH', 'Chaume'), ('BO', 'Bois'), ('PI', 'Pierre')]
POSTES = [('structure', 'Structure des plans (sans la colonnade)'), ('cultures', 'Carrés de culture (Plant X)'),
          ('ancres_temporaires', "Ancres temporaires (démolies)"), ('pics', 'Murs de pics'),
          ('tourelles_chaume', 'Murs de chaume (démolis)'), ('equipements', 'Équipements'),
          ('tourelles', 'Tourelles (déjà en inventaire)')]
NOMS = dict(fz.NOMS, demi_pilier='Demi-pilier métal (ancre)')
F_ = 'Arial'
base = Font(name=F_, size=10); bold = Font(name=F_, size=10, bold=True)
blue = Font(name=F_, size=10, color='0000FF'); green = Font(name=F_, size=10, color='008000')
title = Font(name=F_, size=14, bold=True)
head_fill = PatternFill('solid', fgColor='DDE4EE'); tot_fill = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); box = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);-'

wb = Workbook()
# ------------------------------------------------------------------ Détail
d = wb.active; d.title = 'Détail'
d['A1'] = 'Bilan matériaux — détail par pièce'; d['A1'].font = title
d['A2'] = ("Bleu = donnée saisie (nombre de pièces, coût unitaire) : modifiez-la, tout se recalcule. Noir = formule. "
           "Coûts unitaires : module 07 §6 et README §9 ; « oui » en colonne Estimation = coût non trouvé dans les sources.")
d['A2'].font = Font(name=F_, size=9, italic=True)
h = 4
cols = ['Poste', 'Pièce', 'Nombre', 'Estimation'] + [f'{n} / u' for c, n in RES] + [f'{n} total' for c, n in RES]
for j, v in enumerate(cols, 1):
    c = d.cell(h, j, v); c.font = bold; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = box
r = h + 1
first = r
for poste, lib in POSTES:
    pieces = S['bilan_complet'][poste]['pieces']
    for cle, n in sorted(pieces.items(), key=lambda kv: -kv[1]):
        cout = fz.COUTS[cle]
        d.cell(r, 1, lib).font = base
        d.cell(r, 2, NOMS[cle]).font = base
        d.cell(r, 3, n).font = blue
        d.cell(r, 4, 'oui' if cout.get('est') else '').font = base
        for k, (code, _) in enumerate(RES):
            u = d.cell(r, 5 + k, cout.get(code, 0)); u.font = blue; u.number_format = NUM
            t = d.cell(r, 5 + len(RES) + k, f'=$C{r}*{L(5 + k)}{r}'); t.font = base; t.number_format = NUM
        for j in range(1, len(cols) + 1):
            d.cell(r, j).border = box
        r += 1
last = r - 1
d.cell(r, 1, 'Total général (avec tourelles)').font = bold
d.cell(r, 3, f'=SUM(C{first}:C{last})').font = bold
for k in range(len(RES)):
    c = d.cell(r, 5 + len(RES) + k, f'=SUM({L(5 + len(RES) + k)}{first}:{L(5 + len(RES) + k)}{last})')
    c.font = bold; c.number_format = NUM
for j in range(1, len(cols) + 1):
    d.cell(r, j).fill = tot_fill; d.cell(r, j).border = box
d.column_dimensions['A'].width = 30; d.column_dimensions['B'].width = 42; d.column_dimensions['C'].width = 9
d.column_dimensions['D'].width = 10
for j in range(5, len(cols) + 1):
    d.column_dimensions[L(j)].width = 11
d.row_dimensions[h].height = 42
d.freeze_panes = d.cell(h + 1, 3)
RNG = lambda col: f"Détail!${col}${first}:${col}${last}"

# ------------------------------------------------------------------ Synthèse
s = wb.create_sheet('Synthèse', 0)
s['A1'] = 'Forteresse Ø30 — bilan matériaux'; s['A1'].font = title
T_ = S['tourelles']['posees']
s['A2'] = (f"Calculé depuis l'onglet Détail. Les tourelles ({T_['heavy']} Heavy, {T_['tek']} Tek, {T_['tesla']} Tesla) "
           f"sont déjà en inventaire ; les {T_['crop']} graines de Plant Species X se plantent à la main. Ne fabriquez "
           "que la ligne « À fabriquer ».")
s['A2'].font = Font(name=F_, size=9, italic=True)
s['A4'] = 'Nombre de structures'; s['A4'].font = bold
hdr = ['Poste', 'Structures', 'Devenir']
for j, v in enumerate(hdr, 1):
    c = s.cell(5, j, v); c.font = bold; c.fill = head_fill; c.border = box
devenir = {'structure': 'définitives', 'cultures': 'définitifs (Plant X à planter)',
           'ancres_temporaires': 'démolies à la fin', 'pics': 'définitifs, posés à la main',
           'tourelles_chaume': 'démolis après la pose des tourelles', 'equipements': 'définitifs, hors plans',
           'tourelles': 'définitives, déjà en inventaire'}
rr = 6
for poste, lib in POSTES:
    s.cell(rr, 1, lib).font = base
    c = s.cell(rr, 2, f'=SUMIFS({RNG("C")},{RNG("A")},A{rr})'); c.font = green; c.number_format = NUM
    s.cell(rr, 3, devenir[poste]).font = base
    for j in range(1, 4): s.cell(rr, j).border = box
    rr += 1
s.cell(rr, 1, 'Total à poser').font = bold
c = s.cell(rr, 2, f'=SUM(B6:B{rr - 1})'); c.font = bold; c.number_format = NUM
s.cell(rr, 3, f'=B{rr}-B8-B10').font = base
s.cell(rr, 3).number_format = '#,##0" restent en place"'
for j in range(1, 4): s.cell(rr, j).fill = tot_fill; s.cell(rr, j).border = box
tot_struct_row = rr

rr += 2
s.cell(rr, 1, 'Matériaux par poste').font = bold
rr += 1
hdr_row = rr
s.cell(rr, 1, 'Poste').font = bold
for k, (code, nom) in enumerate(RES):
    s.cell(rr, 2 + k, nom).font = bold
for j in range(1, 2 + len(RES)):
    s.cell(rr, j).fill = head_fill; s.cell(rr, j).border = box; s.cell(rr, j).alignment = Alignment(wrap_text=True, vertical='center')
s.row_dimensions[rr].height = 30
rr += 1
p0 = rr
for poste, lib in POSTES:
    s.cell(rr, 1, lib).font = base
    for k in range(len(RES)):
        col = L(5 + len(RES) + k)
        c = s.cell(rr, 2 + k, f'=SUMIFS({RNG(col)},{RNG("A")},$A{rr})'); c.font = green; c.number_format = NUM
    for j in range(1, 2 + len(RES)): s.cell(rr, j).border = box
    rr += 1
p_tour = rr - 1
lines = [('À fabriquer (hors tourelles)', lambda col: f'=SUM({col}{p0}:{col}{p_tour - 1})'),
         ('Total avec les tourelles', lambda col: f'=SUM({col}{p0}:{col}{p_tour})')]
tot_rows = {}
for lib, f in lines:
    s.cell(rr, 1, lib).font = bold
    for k in range(len(RES)):
        c = s.cell(rr, 2 + k, f(L(2 + k))); c.font = bold; c.number_format = NUM
    for j in range(1, 2 + len(RES)): s.cell(rr, j).fill = tot_fill; s.cell(rr, j).border = box
    tot_rows[lib] = rr
    rr += 1
fab = tot_rows['À fabriquer (hors tourelles)']; avec = tot_rows['Total avec les tourelles']

# ressources brutes
rr += 1
s.cell(rr, 1, 'Équivalent en ressources brutes').font = bold
rr += 1
s.cell(rr, 1, 'Recette (hypothèse, à vérifier sur votre serveur)').font = bold
rr += 1
hyp = {}
for lib, v, key in [('Minerai par lingot', 2, 'min'), ('Perles de silice par électronique', 3, 'sil'),
                    ('Lingot par électronique', 1, 'lel'), ('Obsidienne par polymère', 2, 'obs'),
                    ('Pâte par polymère', 2, 'ppo'), ('Chitine/kératine par pâte', 4, 'chi'), ('Pierre par pâte', 8, 'pie'),
                    ('Silex par poudre étincelante', 2, 'sx'), ('Pierre par poudre étincelante', 1, 'ppe')]:
    s.cell(rr, 1, lib).font = base
    c = s.cell(rr, 2, v); c.font = blue; c.fill = PatternFill('solid', fgColor='FFFF00')
    hyp[key] = f'$B${rr}'
    rr += 1
rr += 1
for j, v in enumerate(['Ressource brute', 'Hors tourelles', 'Avec les tourelles'], 1):
    c = s.cell(rr, j, v); c.font = bold; c.fill = head_fill; c.border = box
rr += 1
col = {code: L(2 + k) for k, (code, _) in enumerate(RES)}
def lingots(row): return f"({col['MI']}{row}+{hyp['lel']}*{col['EL']}{row})"
def pate(row): return f"({col['PC']}{row}+{hyp['ppo']}*{col['PO']}{row})"
brut = [('Lingots de métal au total (électroniques comprises)', lambda R: f'={lingots(R)}'),
        ('Minerai de métal', lambda R: f"={hyp['min']}*{lingots(R)}"),
        ('Pâte cimentée au total (polymère compris)', lambda R: f'={pate(R)}'),
        ('Chitine ou kératine', lambda R: f"={hyp['chi']}*{pate(R)}"),
        ('Pierre', lambda R: f"={hyp['pie']}*{pate(R)}+{hyp['ppe']}*{col['PE']}{R}+{col['PI']}{R}"),
        ('Obsidienne', lambda R: f"={hyp['obs']}*{col['PO']}{R}"),
        ('Perles de silice', lambda R: f"={hyp['sil']}*{col['EL']}{R}"),
        ('Silex', lambda R: f"={hyp['sx']}*{col['PE']}{R}")]
for lib, f in brut:
    s.cell(rr, 1, lib).font = base
    for j, R in ((2, fab), (3, avec)):
        c = s.cell(rr, j, f(R)); c.font = base; c.number_format = NUM
    for j in range(1, 4): s.cell(rr, j).border = box
    rr += 1
rr += 1
col_ = S['options']['colonnade']['bilan']
reh_ = S['options']['rehausse_par_niveau']
notes = [f"Option colonnade (plan 2, {S['options']['colonnade']['total']} piliers) : +{col_['MI']:,} lingots et "
         f"+{col_['PC']:,} pâte, non comptés ici.".replace(',', ' '),
         f"Option rehausse du pod, par niveau ({reh_['total']} pièces) : +{reh_['bilan']['MI']:,} lingots et "
         f"+{reh_['bilan']['PC']:,} pâte.".replace(',', ' '),
         "Étoile centrale déjà construite : retirez 87 fondations et 174 triangles (−8 700 lingots, −2 697 pâte) dans l'onglet Détail.",
         "Source des quantités : outils/forteresse.py → donnees/synthese.json (bilan_complet). Limites moteur : module 09."]
for n in notes:
    s.cell(rr, 1, n).font = Font(name=F_, size=9, italic=True); rr += 1
s.column_dimensions['A'].width = 46
for k in range(len(RES)): s.column_dimensions[L(2 + k)].width = 13
s.column_dimensions['C'].width = 22

# ------------------------------------------------------------------ Par plan
pp = wb.create_sheet('Par plan')
pp['A1'] = 'Structure par plan (hors tourelles, carrés de culture et murs de chaume)'; pp['A1'].font = title
pp['A2'] = ("Nombre de pièces par plan (bleu) ; coûts unitaires repris de l'onglet Détail (vert). Le plan 2 (colonnade) "
            "est une option, hors du total des autres onglets.")
pp['A2'].font = Font(name=F_, size=9, italic=True)
cles = sorted({k for v in S['plans'].values() for k in v['pieces'] if k not in ('heavy', 'tek', 'tesla', 'chaume_mur', 'crop')},
              key=lambda k: -S['bilan_complet']['structure']['pieces'].get(k, 0))
hr = 6
pp.cell(4, 1, 'Lingots / pièce').font = bold; pp.cell(5, 1, 'Pâte / pièce').font = bold
pp.cell(hr, 1, 'Plan').font = bold; pp.cell(hr, 1).fill = head_fill
for j, k in enumerate(cles, 2):
    c = pp.cell(hr, j, NOMS[k]); c.font = bold; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = box
    for rrow, code in ((4, 'MI'), (5, 'PC')):
        ci = 5 + [c_ for c_, _ in RES].index(code)
        if any(k in S['bilan_complet'][po]['pieces'] for po, _ in POSTES):
            f = (f'=INDEX(Détail!${L(ci)}${first}:${L(ci)}${last},MATCH("{NOMS[k]}",Détail!$B${first}:$B${last},0))')
            x = pp.cell(rrow, j, f); x.font = green
        else:                                    # pièce présente seulement dans une option (colonnade)
            x = pp.cell(rrow, j, fz.COUTS[k].get(code, 0)); x.font = blue
        x.number_format = NUM
nj = len(cles)
for t, lab in enumerate(['Pièces', 'Lingots', 'Pâte']):
    c = pp.cell(hr, 2 + nj + t, lab); c.font = bold; c.fill = head_fill; c.border = box
pp.row_dimensions[hr].height = 54
r = hr + 1; r0 = r
for plan, v in S['plans'].items():
    pp.cell(r, 1, plan.replace('_', ' ')).font = base
    for j, k in enumerate(cles, 2):
        c = pp.cell(r, j, v['pieces'].get(k, 0)); c.font = blue; c.number_format = NUM; c.border = box
    a, b = L(2), L(1 + nj)
    pp.cell(r, 2 + nj, f'=SUM({a}{r}:{b}{r})').number_format = NUM
    pp.cell(r, 3 + nj, f'=SUMPRODUCT({a}{r}:{b}{r},{a}$4:{b}$4)').number_format = NUM
    pp.cell(r, 4 + nj, f'=SUMPRODUCT({a}{r}:{b}{r},{a}$5:{b}$5)').number_format = NUM
    r += 1
pp.cell(r, 1, 'Total').font = bold
for j in range(2, 5 + nj):
    c = pp.cell(r, j, f'=SUM({L(j)}{r0}:{L(j)}{r - 1})'); c.font = bold; c.number_format = NUM; c.fill = tot_fill
pp.cell(r + 2, 1, "Murs de chaume (dans les plans de zone), tourelles et carrés de culture (plans 7, 8 et 11) : onglet Détail.").font = Font(name=F_, size=9, italic=True)
pp.column_dimensions['A'].width = 40
for j in range(2, 5 + nj): pp.column_dimensions[L(j)].width = 11
pp.freeze_panes = 'B7'

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F_:
                c.font = Font(name=F_, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
out = os.path.join(fz.RACINE, 'donnees', '07_bilan_materiaux.xlsx')
wb.save(out)
print(out)
